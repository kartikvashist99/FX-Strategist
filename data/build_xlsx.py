from datetime import date, datetime
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
import export as X
import gen_data as g

F = "Arial"
NAVY = "0B2A4A"
HDR = PatternFill("solid", fgColor=NAVY)
HFONT = Font(name=F, bold=True, color="FFFFFF", size=10)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=16, bold=True, color=NAVY)
SUB = Font(name=F, size=10, italic=True, color="666666")
BLUE = Font(name=F, size=10, color="0000FF")
BANDFILL = {"High": "D9F2E1", "Medium": "FDEBD0", "Low": "F8D7DA", "Locked": "E2E6EC"}
BANNER = PatternFill("solid", fgColor="FFF4CC")
thin = Side(style="thin", color="D0D5DD")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

INRF = '#,##0;(#,##0);"-"'
USDF = '#,##0.00;(#,##0.00);"-"'
DATEF = "dd-mmm-yyyy"

wb = Workbook()

def dt(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, date) and not isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    return v

def table(ws, r0, headers, rows, fmts=None, widths=None):
    fmts = fmts or {}
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h); c.font = HFONT; c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + i, j, dt(v)); c.font = BODY
            if j in fmts:
                c.number_format = fmts[j]
            elif isinstance(c.value, datetime):
                c.number_format = DATEF
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(rows)

def data_sheet(name, df, cols, headers, fmts, widths):
    ws = wb.create_sheet(name)
    rows = [[r[c] for c in cols] for r in df.to_dict("records")]
    table(ws, 1, headers, rows, fmts, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
    return ws

# ---------------- README ----------------
ws = wb.active; ws.title = "README"
ws["A1"] = "Lumio Analytics - Synthetic Treasury Dataset"; ws["A1"].font = TITLE
ws["A2"] = "Built for the Razorpay AI PM Build Challenge prototype (EEFC conversion-planning agent)"; ws["A2"].font = SUB
ws["A4"] = "ALL DATA IN THIS WORKBOOK IS SYNTHETIC. Company, clients, vendors and the USD/INR rate path are fictional; rules and cost assumptions are listed below."
ws["A4"].font = BOLD; ws["A4"].fill = BANNER
ws.merge_cells("A4:D4"); ws.row_dimensions[4].height = 30; ws["A4"].alignment = Alignment(wrap_text=True, vertical="center")
r = table(ws, 6, ["Assumption", "Value"], list(g.ASSUMPTIONS.items()))
r += 2
ws.cell(r, 1, "Story in one line").font = BOLD
ws.cell(r + 1, 1, "For 22 months Lumio converted every USD receipt to INR at settlement, then bought USD back to pay AWS, tools and a US contractor. "
        "On 1 Sep 2026 it switched on EEFC settlement and now holds USD, but manages conversions by hand (one manual conversion on 29 Sep). "
        "On 26 Oct 2026 ('today') it has USD in two EEFC buckets, an INR balance sitting at its buffer, payroll in 7 days and a September deadline in 5.").alignment = Alignment(wrap_text=True)
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=4); ws.row_dimensions[r + 1].height = 60
r += 3
guide = [
    ("Demo_Scenario", "Today's position, the agent's 6 expected outputs (locked / green / amber / red) and the use-case coverage checklist"),
    ("Value_at_Stake", "Baseline (convert at settlement) vs netting from EEFC - live formulas"),
    ("Monthly_Summary", "24-month roll-up of receipts, conversions, outflows and balances - live formulas over the data sheets"),
    ("Detected_Patterns", "Recurring obligations the agent infers from history, with confidence and its components"),
    ("Client_Behaviour", "Payment timing per client (on-time, late payer, new client, churned, annual)"),
    ("Receipts", "Every USD invoice: issued, due, paid / outstanding / scheduled"),
    ("INR_Outflows", "Payroll, rent, GST, advance tax, fees, one-offs, FD sweeps (non-operating)"),
    ("USD_Outflows", "AWS, tools, contractor - how each was paid and the avoidable FX cost"),
    ("Conversions", "Every USD to INR conversion (auto at settlement, then manual from EEFC)"),
    ("EEFC_Ledger", "Credits and debits by monthly bucket since EEFC was switched on"),
    ("Daily_Balances", "INR current-account and EEFC USD balance, every day"),
    ("FX_Rates", "Synthetic USD/INR mid rate, business days"),
    ("Clients", "Client master"),
]
r = table(ws, r, ["Sheet", "What it holds"], guide)
ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 110
for row in ws.iter_rows(min_row=7, max_row=r):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")

# ---------------- Demo_Scenario ----------------
D = X.data
ws = wb.create_sheet("Demo_Scenario")
ws["A1"] = "Scenario: Monday 26 Oct 2026"; ws["A1"].font = TITLE
ws["A2"] = "Agent's expected outputs for the prototype. Confidence = forecast confidence in the cash inputs, not a rate prediction."; ws["A2"].font = SUB
pos = D["position_today"]
ws["A4"] = "Position today"; ws["A4"].font = BOLD
rows = [("INR current account (Rs)", pos["inr_balance"], "At the Rs 8L buffer policy"),
        ("USD/INR today", D["meta"]["usd_inr_today"], "Synthetic"),
        ("EEFC Sep-2026 bucket (USD)", pos["eefc_buckets"][0]["usd"], "Deadline 31 Oct 2026 - 5 days left"),
        ("EEFC Oct-2026 bucket (USD)", pos["eefc_buckets"][1]["usd"], "Deadline 30 Nov 2026 - 35 days left"),
        ("Next payroll", "Mon 2 Nov 2026", f"~Rs {X.pr['estimate']/1e5:.1f}L")]
table(ws, 5, ["Item", "Value", "Note"], rows)
for rr in (6,):
    ws.cell(rr, 2).number_format = INRF
ws.cell(7, 2).number_format = "0.00"; ws.cell(8, 2).number_format = USDF; ws.cell(9, 2).number_format = USDF
r = 12
ws.cell(r, 1, "Agent plan (expected outputs)").font = BOLD
cols = ["id", "type", "title", "action", "usd", "est_inr", "deadline", "confidence", "band", "reasoning", "flagged_input", "merchant_can"]
hdrs = ["ID", "Type", "Title", "Action", "USD", "Est. INR", "Deadline / timing", "Confidence %", "Band", "Agent reasoning", "Input flagged for review", "What the merchant can do"]
recs = [[x[c] for c in cols] for x in D["recommendations"]]
end = table(ws, r + 1, hdrs, recs, {5: USDF, 6: INRF})
for i in range(r + 2, end + 1):
    b = ws.cell(i, 9).value
    for j in range(1, len(cols) + 1):
        ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
    if b in BANDFILL:
        ws.cell(i, 9).fill = PatternFill("solid", fgColor=BANDFILL[b]); ws.cell(i, 8).fill = PatternFill("solid", fgColor=BANDFILL[b])
for j, w in enumerate([6, 22, 30, 40, 12, 13, 22, 12, 10, 60, 32, 34], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
r = end + 2
ws.cell(r, 1, "Bands").font = BOLD
table(ws, r + 1, ["Band", "Rule", "Behaviour"], [
    ("High", ">= 80%", "Quick approve: one summary card, then app authentication"),
    ("Medium", "50-79%", "Flagged input shown; merchant confirms/edits it, then authenticates"),
    ("Low", "< 50%", "No recommendation - the agent asks a question first"),
    ("Locked", "Rule-driven", "EEFC deadline conversions; no score; cannot be skipped")])
for i, bnd in enumerate(["High", "Medium", "Low", "Locked"], r + 2):
    ws.cell(i, 1).fill = PatternFill("solid", fgColor=BANDFILL[bnd])
r = r + 8
ws.cell(r, 1, "Use-case coverage checklist").font = BOLD
table(ws, r + 1, ["Case", "Where it appears"], X.coverage)
ws.freeze_panes = "A4"

# ---------------- Value_at_Stake (formulas) ----------------
ws = wb.create_sheet("Value_at_Stake")
ws["A1"] = "Value at stake - netting USD costs from EEFC"; ws["A1"].font = TITLE
ws["A2"] = "Formulas read the USD_Outflows sheet. Blue = assumption you can change on README / generator."; ws["A2"].font = SUB
items = [
    ("EEFC switched on (date)", datetime(2026, 9, 1), DATEF, "Input"),
    ("Months of convert-at-settlement history", 22, "0", "Input"),
    ("USD costs paid by buying USD back (USD)", '=SUMIFS(USD_Outflows!$D:$D,USD_Outflows!$A:$A,"<"&B4)', USDF, ""),
    ("INR spent on those USD costs (Rs)", '=SUMIFS(USD_Outflows!$I:$I,USD_Outflows!$A:$A,"<"&B4)', INRF, ""),
    ("Avoidable FX cost - markups + wire fees (Rs)", '=SUMIFS(USD_Outflows!$J:$J,USD_Outflows!$A:$A,"<"&B4)', INRF, ""),
    ("Avoidable FX cost, annualised (Rs / yr)", "=B8/B5*12", INRF, ""),
    ("As % of USD costs", "=B8/B7", "0.0%", ""),
    ("Since EEFC: USD costs paid directly from EEFC (USD)", '=SUMIFS(USD_Outflows!$D:$D,USD_Outflows!$A:$A,">="&B4,USD_Outflows!$G:$G,"EEFC (USD)")', USDF, ""),
    ("Since EEFC: still paid by card (EEFC empty on 3 Sep) (USD)", '=SUMIFS(USD_Outflows!$D:$D,USD_Outflows!$A:$A,">="&B4,USD_Outflows!$G:$G,"INR (bought USD)")', USDF, ""),
]
table(ws, 3, ["Metric", "Value", "Note"], [])
for i, (lab, val, fmt, note) in enumerate(items, 4):
    ws.cell(i, 1, lab).font = BODY
    c = ws.cell(i, 2, val); c.number_format = fmt; c.font = BLUE if note == "Input" else BODY
    ws.cell(i, 3, note).font = SUB
ws.cell(14, 1, "Not claimed: any gain from predicting USD/INR moves. The agent's value is removing avoidable costs (round-trips, deadline-day forced conversions) and timing conversions to actual INR needs.").font = SUB
ws.merge_cells("A14:C14"); ws.row_dimensions[14].height = 30; ws["A14"].alignment = Alignment(wrap_text=True)
ws.column_dimensions["A"].width = 58; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 30

# ---------------- Data sheets ----------------
inv = X.INV
data_sheet("Receipts", inv, ["invoice_id", "client_id", "client", "plan", "currency", "amount_usd", "invoice_date", "due_date", "paid_date", "days_after_due", "status"],
           ["Invoice ID", "Client ID", "Client", "Plan", "Currency", "Amount (USD)", "Invoice date", "Due date", "Paid date", "Days after due", "Status"],
           {6: USDF, 10: "0"}, [20, 9, 26, 9, 9, 13, 13, 13, 13, 12, 26])
data_sheet("INR_Outflows", X.INR_HIST, ["date", "category", "payee", "amount_inr", "recurring"],
           ["Date", "Category", "Payee", "Amount (Rs)", "Recurring?"], {4: INRF}, [13, 24, 34, 15, 26])
data_sheet("USD_Outflows", X.USDP, ["date", "category", "payee", "amount_usd", "channel", "recurring", "paid_from", "mid_rate", "inr_cost", "avoidable_fx_cost_inr"],
           ["Date", "Category", "Payee", "Amount (USD)", "Channel", "Recurring?", "Paid from", "Mid rate", "INR cost (Rs)", "Avoidable FX cost (Rs)"],
           {4: USDF, 8: "0.0000", 9: INRF, 10: INRF}, [13, 16, 30, 13, 15, 11, 17, 10, 14, 16])
data_sheet("Conversions", X.CONV, ["date", "type", "usd", "rate", "inr", "source", "from_bucket"],
           ["Date", "Type", "USD", "Rate", "INR (Rs)", "Source", "From EEFC bucket"], {3: USDF, 4: "0.0000", 5: INRF}, [13, 36, 13, 10, 15, 20, 16])
data_sheet("EEFC_Ledger", X.LEDGER, ["date", "entry", "ref", "usd", "bucket"],
           ["Date", "Entry", "Reference", "USD (+credit / -debit)", "Bucket (month received)"], {4: USDF}, [13, 40, 36, 18, 20])
data_sheet("Daily_Balances", X.DAILY, ["date", "inr_balance", "eefc_usd"], ["Date", "INR balance (Rs)", "EEFC balance (USD)"], {2: INRF, 3: USDF}, [13, 17, 17])
import pandas as pd
fx = pd.DataFrame([dict(date=d, usd_inr=v) for d, v in g.FX.items() if d >= g.START])
data_sheet("FX_Rates", fx, ["date", "usd_inr"], ["Date", "USD/INR mid (synthetic)"], {2: "0.0000"}, [13, 20])
cl = pd.DataFrame([dict(id=c["id"], name=c["name"], country=c["country"], plan="Annual" if c["cycle"] == "annual" else "Monthly",
                        list_price=round(c["amt"] * g.SCALE, -1), start=f"{c['start'][0]}-{c['start'][1]:02d}",
                        end=(f"{c['end'][0]}-{c['end'][1]:02d}" if c["end"] else ""), behaviour={"ontime": "Pays on time", "late": "Late payer (30-45 days past due)", "new_large": "New large client (Aug 2026)"}[c["beh"]]) for c in g.CLIENTS])
data_sheet("Clients", cl, ["id", "name", "country", "plan", "list_price", "start", "end", "behaviour"],
           ["ID", "Client", "Country", "Plan", "Price at start (USD)", "Start", "End", "Behaviour"], {5: USDF}, [7, 28, 10, 9, 16, 9, 9, 34])

# ---------------- Detected patterns / client behaviour ----------------
P = X.PAT
ws = data_sheet("Detected_Patterns", P, ["obligation", "currency", "cadence", "occurrences_seen", "last_seen", "next_expected", "estimate", "range_low", "range_high",
                                         "confidence", "band", "history_factor", "timing_factor", "amount_factor", "trend_break", "agent_note"],
                ["Obligation", "Ccy", "Cadence", "Times seen", "Last seen", "Next expected", "Estimate", "Range low", "Range high", "Confidence %", "Band",
                 "History factor", "Timing factor", "Amount factor", "Trend break?", "Agent note"],
                {7: USDF, 8: USDF, 9: USDF, 12: "0.00", 13: "0.00", 14: "0.00"}, [28, 6, 18, 10, 13, 13, 13, 12, 12, 12, 10, 10, 10, 10, 10, 70])
for i in range(2, len(P) + 2):
    for j in (5, 6):
        v = ws.cell(i, j).value
        if isinstance(v, str):
            ws.cell(i, j).value = datetime.fromisoformat(v); ws.cell(i, j).number_format = DATEF
    b = ws.cell(i, 11).value
    ws.cell(i, 11).fill = PatternFill("solid", fgColor=BANDFILL.get(b, "FFFFFF"))
n = len(P) + 3
ws.cell(n, 1, "Confidence = 100 x history factor x timing factor x amount factor. History = sqrt(min(1,(n+3)/12)); timing = 1 - min(0.4, sd of day/15) (0.6 if seen once); "
        "amount = 1 - min(0.5, CV of last 6), x0.8 if the latest value breaks >10% from the prior average. Illustrative heuristic for the prototype.").font = SUB
ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=12); ws.row_dimensions[n].height = 30
ws.cell(n, 1).alignment = Alignment(wrap_text=True)

R = X.RECV
ws = data_sheet("Client_Behaviour", R, ["client_id", "client", "plan", "status", "invoices_paid", "avg_days_after_due", "sd_days", "latest_amount_usd", "timing_confidence", "band"],
                ["ID", "Client", "Plan", "Status", "Invoices paid", "Avg days after due (- = early)", "SD (days)", "Latest invoice (USD)", "Timing confidence %", "Band"],
                {6: "0.0", 7: "0.0", 8: USDF}, [7, 28, 9, 30, 12, 16, 10, 16, 14, 10])
for i in range(2, len(R) + 2):
    b = ws.cell(i, 10).value
    ws.cell(i, 10).fill = PatternFill("solid", fgColor=BANDFILL.get(b, "FFFFFF"))

# ---------------- Monthly_Summary (formulas) ----------------
ws = wb.create_sheet("Monthly_Summary", 3)
ws["A1"] = "Monthly summary (live formulas over the data sheets)"; ws["A1"].font = TITLE
hdr = ["Month", "Month end", "USD received", "INR from conversions (Rs)", "INR operating outflows (Rs)", "FD sweeps (Rs)", "USD costs (USD)", "Avoidable FX cost (Rs)", "INR balance at month end (Rs)", "EEFC USD at month end"]
table(ws, 3, hdr, [])
months = g.HIST_MONTHS
for i, (y, m) in enumerate(months, 4):
    ws.cell(i, 1, datetime(y, m, 1)).number_format = "mmm-yyyy"
    ws.cell(i, 2, f"=MIN(EOMONTH(A{i},0),DATE(2026,10,26))").number_format = DATEF
    ws.cell(i, 3, f'=SUMIFS(Receipts!$F:$F,Receipts!$I:$I,">="&A{i},Receipts!$I:$I,"<="&B{i})').number_format = USDF
    ws.cell(i, 4, f'=SUMIFS(Conversions!$E:$E,Conversions!$A:$A,">="&A{i},Conversions!$A:$A,"<="&B{i})').number_format = INRF
    ws.cell(i, 5, f'=SUMIFS(INR_Outflows!$D:$D,INR_Outflows!$A:$A,">="&A{i},INR_Outflows!$A:$A,"<="&B{i},INR_Outflows!$B:$B,"<>Sweep to fixed deposit")').number_format = INRF
    ws.cell(i, 6, f'=SUMIFS(INR_Outflows!$D:$D,INR_Outflows!$A:$A,">="&A{i},INR_Outflows!$A:$A,"<="&B{i},INR_Outflows!$B:$B,"Sweep to fixed deposit")').number_format = INRF
    ws.cell(i, 7, f'=SUMIFS(USD_Outflows!$D:$D,USD_Outflows!$A:$A,">="&A{i},USD_Outflows!$A:$A,"<="&B{i})').number_format = USDF
    ws.cell(i, 8, f'=SUMIFS(USD_Outflows!$J:$J,USD_Outflows!$A:$A,">="&A{i},USD_Outflows!$A:$A,"<="&B{i})').number_format = INRF
    ws.cell(i, 9, f"=INDEX(Daily_Balances!$B:$B,MATCH(B{i},Daily_Balances!$A:$A,0))").number_format = INRF
    ws.cell(i, 10, f"=INDEX(Daily_Balances!$C:$C,MATCH(B{i},Daily_Balances!$A:$A,0))").number_format = USDF
    for j in range(1, 11):
        ws.cell(i, j).font = BODY
t = len(months) + 4
ws.cell(t, 1, "Total").font = BOLD
for j, col in zip(range(3, 9), "CDEFGH"):
    c = ws.cell(t, j, f"=SUM({col}4:{col}{t-1})"); c.font = BOLD; c.number_format = USDF if col in "CG" else INRF
for j, w in enumerate([11, 13, 14, 18, 18, 14, 14, 16, 18, 16], 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "B4"

order = ["README", "Demo_Scenario", "Value_at_Stake", "Monthly_Summary", "Detected_Patterns", "Client_Behaviour", "Receipts", "INR_Outflows",
         "USD_Outflows", "Conversions", "EEFC_Ledger", "Daily_Balances", "FX_Rates", "Clients"]
wb._sheets = [wb[s] for s in order]
wb.save("Lumio_Synthetic_Treasury_Data.xlsx")
print("saved")
